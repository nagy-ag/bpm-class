"""Explicit readiness probes for the user's imported nap01/03/07/08.

Not coursework completion or UiPath execution. Never executes lesson code/HTML.
Uses fresh isolated data directories; outputs JSON/CSV diagnostics, not submissions.
Run with bundled Python: -m harness.course_tools.readiness_updated PHASE --out PATH.
"""
import argparse
import ast
import csv
import hashlib
import json
import math
import re
import shutil
import sys
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from harness.archives import extract_zip

ROOT = Path(__file__).resolve().parents[2]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')


def csv_roundtrip(path, header, rows):
    with path.open('x', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)
    with path.open(encoding='utf-8-sig', newline='') as stream:
        actual = list(csv.reader(stream))
    assert actual == [header] + [[str(v) for v in row] for row in rows]


def local03(out):
    from lxml import etree
    ns = {'b': 'http://www.omg.org/spec/BPMN/20100524/MODEL',
          'd': 'http://www.omg.org/spec/BPMN/20100524/DI'}
    tree = etree.parse(str(ROOT/'nap01/nap01/docs/kave_rendeles_mini.bpmn'))
    nodes = tree.xpath('//b:process/*[not(self::b:sequenceFlow)]', namespaces=ns)
    flows = tree.xpath('//b:sequenceFlow', namespaces=ns)
    ids = {n.get('id') for n in nodes}
    assert len(nodes) == 5 and len(flows) == 4
    assert all(f.get('sourceRef') in ids and f.get('targetRef') in ids for f in flows)
    assert len(tree.xpath('//d:BPMNShape', namespaces=ns)) == 5
    assert len(tree.xpath('//d:BPMNEdge', namespaces=ns)) == 4
    assert [(f.get('sourceRef'), f.get('targetRef')) for f in flows] == [('n0','n1'),('n1','n2'),('n2','n3'),('n3','n4')]
    fixture = json.loads((ROOT/'nap03/nap03/docs/cmc_pelda_valasz.json').read_text())
    assert fixture['status']['error_code'] == 0
    for coin in ('1','1027'):
        q = fixture['data'][coin]['quote']['EUR']
        assert q['price'] > 0 and math.isfinite(q['percent_change_24h'])
    assert fixture['data']['1']['quote']['EUR']['percent_change_24h'] == -1.37
    for path in (ROOT/'nap03/nap03/docs').glob('*colab.py'):
        ast.parse(path.read_text(encoding='utf-8-sig'))
    assert not (ROOT/'nap03/nap03/pages/06_mini_lanc.html').exists()
    return {'cafe_nodes':5,'connected_flows':4,'diagram_geometry':True,
            'fixture_percent_points':-1.37,'colab_scripts_syntax':'pass',
            'removed_chain_excluded':True,'bpmn_browser_or_colab_execution':False}


def live(out):
    if (out/'crm.json').exists():
        raise FileExistsError('Existing CRM checkpoint requires reconciliation; use a fresh output directory.')
    from harness.course_tools.day03.bpa_access import CourseAPI, AccessError
    api = CourseAPI()
    quotes = api.cmc_quotes((1,1027,5426),'EUR')
    assert quotes['status']['error_code'] == 0
    clean = []
    for cid, symbol in ((1,'BTC'),(1027,'ETH'),(5426,'SOL')):
        item = quotes['data'][str(cid)]
        q = item['quote']['EUR']
        assert item['symbol'] == symbol and q['price'] > 0 and math.isfinite(q['price'])
        assert math.isfinite(q['percent_change_24h']) and q['last_updated']
        clean.append({'id':cid,'symbol':symbol,'price_eur':q['price'],
                      'percent_points':q['percent_change_24h'],'provider_time':q['last_updated']})
    save(out/'quotes.json', {'checked_utc':datetime.now(timezone.utc).isoformat(),'rows':clean})
    # Test only this run's unique disposable fictional lead. Never retry uncertain writes.
    state = {'marker':'BPA_READINESS_'+uuid4().hex[:12], 'completed':[], 'cleanup':'not_needed'}
    def checkpoint(): save(out/'crm.json', state)
    def success(response):
        rows = response.get('data',[])
        assert len(rows)==1 and rows[0].get('status')=='success' and rows[0].get('code')=='SUCCESS'
        return rows[0]['details']
    lead_id = None
    checkpoint()
    for module, field in [('Leads','Last_Name'),('Accounts','Account_Name'),('Contacts','Last_Name'),('Deals','Deal_Name')]:
        api.zoho('GET',module,params={'fields':field,'per_page':1})
        state['completed'].append(module+'_read')
        checkpoint()
    try:
        marker = state['marker']
        payload = {'Last_Name':marker,'First_Name':'Fictional readiness', 'Company':marker+' fictional company',
                   'Email':marker.lower()+'@example.invalid','Email_Opt_Out':True,
                   'Description':'Disposable fictional BPA capability test. No communication.'}
        options = {'trigger':[], 'skip_feature_execution':[{'name':'cadences'}]}
        state['cleanup']='write_attempted_reconcile_if_unknown'
        checkpoint()
        lead_id = str(success(api.zoho('POST','Leads',payload={'data':[payload],**options}))['id'])
        state.update(test_lead_id=lead_id,cleanup='pending')
        checkpoint()
        readback = api.zoho('GET','Leads/'+lead_id)['data'][0]
        assert all(readback.get(k)==v for k,v in payload.items())
        state['completed'].append('create_exact_readback')
        checkpoint()
        description = marker+' exact update verified'
        success(api.zoho('PUT','Leads/'+lead_id,payload={'data':[{'id':lead_id,'Description':description}],**options}))
        assert api.zoho('GET','Leads/'+lead_id)['data'][0]['Description']==description
        state['completed'].append('update_exact_readback')
        checkpoint()
    finally:
        if lead_id:
            owned = api.zoho('GET','Leads/'+lead_id)['data'][0]
            assert owned['Last_Name']==state['marker'], 'Ownership mismatch; stop cleanup'
            success(api.zoho('DELETE','Leads/'+lead_id,params={'wf_trigger':'false'}))
            try:
                removed = not api.zoho('GET','Leads/'+lead_id).get('data')
            except AccessError as exc:
                removed = getattr(exc,'status',None) in (400,404)
            assert removed, 'Deletion not verified'
            state['cleanup']='deleted_from_active_crm'
            state['completed'].append('cleanup_verified')
            checkpoint()
    return {'live_quotes':3,'crm':state,'native_excel_or_colab_run':False}


def rows(path, sheet):
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True,data_only=True)
    values=list(book[sheet].values)
    book.close()
    return values


def move_exclusive(source, target):
    """Windows-like Overwrite=False for isolated copies; source retained on collision."""
    with target.open('xb') as stream:
        stream.write(source.read_bytes())
    source.unlink()


def local07(out):
    from pptx import Presentation
    source=ROOT/'nap07/nap07/docs'
    isolated=out/'nap07'
    isolated.mkdir()
    for name in ('ugyfelek.xlsx','Raw_data.xlsx','Dickinson_Sample_Slides.pptx','karbejelentesek_input.xlsx'):
        shutil.copyfile(source/name,isolated/name)
        assert digest(source/name)==digest(isolated/name)
    clients=rows(isolated/'ugyfelek.xlsx','Ugyfelek')
    assert list(clients[0])==['Nev','Varos','EvesDij','Minosites'] and len(clients)==13
    assert all(r[3] is None for r in clients[1:])
    counts={n:sum(r[2]>=n for r in clients[1:]) for n in (100000,150000)}
    assert counts=={100000:6,150000:3}
    assert [v>=100000 for v in (99999,100000,100001)]==[False,True,True]
    assert [v>=150000 for v in (149999,150000,150001)]==[False,True,True]
    raw=rows(isolated/'Raw_data.xlsx','us-500')
    assert list(raw[0])==['State','Quantity','UnitPrice','Total'] and len(raw)==244
    assert all(isinstance(r[1],(int,float)) for r in raw[1:])
    assert len({r[0] for r in raw[1:]})<243, 'Do not aggregate repeated states'
    prs=Presentation(isolated/'Dickinson_Sample_Slides.pptx')
    assert len(prs.slides)==9
    target=[s for s in prs.slides[7].shapes if s.name=='Content Placeholder 6']
    assert len(target)==1
    slide_shapes=[{'name':s.name,'type':str(s.shape_type)} for s in prs.slides[7].shapes]
    csv_roundtrip(isolated/'chart243.csv',['State','Quantity'],[list(r[:2]) for r in raw[1:]])
    csv_roundtrip(isolated/'chart12.csv',['State','Quantity'],[list(r[:2]) for r in raw[1:13]])
    portal=rows(isolated/'karbejelentesek_input.xlsx','bejelentesek')
    assert len(portal)==11 and all(len(r)==6 and r[-1] is None for r in portal[1:])
    folder=extract_zip(source/'gyakorlo_mappa.zip',isolated/'sort')
    files=list(folder.rglob('*'))
    files=[p for p in files if p.is_file()]
    assert len(files)==20
    original={p.name:digest(p) for p in files}
    work=files[0].parent
    mapping={'.pdf':'PDF','.xlsx':'Excel','.docx':'Word','.png':'Kepek'}
    def sort_once():
        count=0
        for p in list(work.iterdir()):
            if p.is_file() and p.suffix.lower() in mapping:
                target=work/mapping[p.suffix.lower()]
                target.mkdir(exist_ok=True)
                move_exclusive(p,target/p.name)
                count+=1
        return count
    assert sort_once()==18 and sort_once()==0
    assert Counter(p.suffix.lower() for p in files)=={'.pdf':6,'.xlsx':4,'.docx':5,'.png':3,'.txt':2}
    assert {p.name:digest(p) for p in folder.rglob('*') if p.is_file()}==original
    (work/'upper.PDF').write_bytes(b'fixture')
    assert sort_once()==1
    (work/'upper.PDF').write_bytes(b'new')
    try: sort_once()
    except FileExistsError: pass
    else: raise AssertionError('Collision overwrote a file')
    assert (work/'upper.PDF').read_bytes()==b'new' and (work/'PDF/upper.PDF').read_bytes()==b'fixture'
    arithmetic={}
    for text,expected in [('20',5000),('5',1250),('0',0),('1000000',250000000),('sok',None),('-1',None),('',None),('1000001',None),('1.5',None),('2147483648',None)]:
        value=int(text) if re.fullmatch(r'[+-]?\d+',text.strip()) else -1
        actual=value*250 if 0<=value<=1000000 else None
        assert actual==expected
        arithmetic[text]=actual
    return {'classification_counts':counts,'raw_rows':243,'mini_rows':12,'slide8_shapes':slide_shapes,
            'sorting':'18 moved, repeat0, uppercase/collision/byte preservation passed','portal_input_rows':10,
            'arithmetic_cases':arithmetic,'native_robot_execution':False,
            'legacy_chart_mismatch':'Old RiportRobot targets slide9 with12/6; current lesson slide8 with243/12.'}


def invoice_fields(path):
    from pypdf import PdfReader
    text='\n'.join(p.extract_text() or '' for p in PdfReader(path).pages)
    if not text.strip(): raise ValueError('business:no_readable_text')
    number=re.search(r'Invoice No:\s*(INV-[0-9-]+)',text)
    amount=re.search(r'Total Due:\s*([0-9,]+\.\d{2})\s*EUR',text)
    if not number or not amount: raise ValueError('business:missing_invoice_field')
    return [number.group(1).strip(),amount.group(1)]


def local08pdf(out):
    work=extract_zip(ROOT/'nap08/nap08/docs/szamlak.zip',out/('invoices-'+uuid4().hex[:8]))
    pdfs=sorted(work.rglob('*.pdf'))
    assert len(pdfs)==11
    pairs=[invoice_fields(p) for p in pdfs]
    assert pairs[0]==['INV-2026-001','3,700.00'] and pairs[-1]==['INV-2026-0011','5,300.00']
    expected=[['INV-2026-00'+str(n),v] for n,v in enumerate(['3,700.00','4,550.00','3,800.00','3,500.00','3,100.00','2,800.00','4,420.00','6,950.00','7,450.00','4,650.00','5,300.00'],1)]
    assert pairs==expected
    from openpyxl import load_workbook
    template=next(work.rglob('*.xlsx'))
    book=load_workbook(template,read_only=True)
    assert 'Sheet1' in book.sheetnames
    book.close()
    # Independent second extraction engine verifies each supplied pair.
    import pdfplumber
    independent=[]
    for path in pdfs:
        with pdfplumber.open(path) as pdf:
            text='\n'.join(p.extract_text() or '' for p in pdf.pages)
        independent.append([re.search(r'Invoice No:\s*(INV-[0-9-]+)',text)[1].strip(),
                            re.search(r'Total Due:\s*([\d,]+\.\d{2})\s*EUR',text)[1]])
    assert independent==pairs
    csv_roundtrip(out/'invoices.csv',['Szamlaszam','Osszeg'],pairs)
    return {'pairs':pairs,'source_directory':str(work),'total_eur_check':str(sum(Decimal(r[1].replace(',','')) for r in pairs)),
            'literal_delimiter_risk':'pypdf returns Total Due colon then newline; pdfplumber puts address on invoice-number line. UiPath text layout must be checked; whitespace-aware number/amount matching is more robust.',
            'two_independent_extractors_agree':True,'output_data_rows':11,'uipath_pdf_run':False}


def subject(raw):
    line=next((s.lstrip() for s in raw.splitlines() if s.lstrip().lower().startswith('subject:')),None)
    # The supplied exercise logs the normalized subject, not its original case.
    return line.split(':',1)[1].strip().lower() if line is not None else '(nincs tárgy)'


def category(title):
    text=title.lower()
    if 'számla' in text or 'szamla' in text: return 'Szamlazas'
    if 'rendelés' in text or 'rendeles' in text: return 'Rendelesek'
    if 'panasz' in text: return 'Panaszok'
    return 'maradt'


def local08email(out):
    work=extract_zip(ROOT/'nap08/nap08/docs/bejovo_emailek.zip',out/('emails-'+uuid4().hex[:8]))
    files=sorted(work.rglob('*.eml'))
    assert len(files)==10
    folder=files[0].parent
    initial={p.name:digest(p) for p in files}
    expected=['Szamlazas','maradt','Szamlazas','Rendelesek','maradt','maradt','maradt','Panaszok','Rendelesek','maradt']
    def process():
        log=[]
        for p in sorted(folder.glob('*.eml')):
            title=subject(p.read_text(encoding='utf-8-sig'))
            dest=category(title)
            log.append([p.name,title,dest])
            if dest!='maradt':
                (folder/dest).mkdir(exist_ok=True)
                move_exclusive(p,folder/dest/p.name)
        return log
    # Save observations before assertion so a course-data discrepancy is inspectable.
    preview=[[p.name,subject(p.read_text(encoding='utf-8-sig')),category(subject(p.read_text(encoding='utf-8-sig')))] for p in files]
    save(out/'email_preview.json',preview)
    actual=[r[2] for r in preview]
    assert actual==expected, actual
    first=process()
    assert Counter(r[2] for r in first)=={'Szamlazas':2,'Rendelesek':2,'Panaszok':1,'maradt':5}
    assert {p.name:digest(p) for p in folder.rglob('*.eml')}==initial
    second=process()
    assert len(second)==5 and all(r[2]=='maradt' for r in second)
    csv_roundtrip(out/'email_first.csv',['Fajl','Targy','Celmappa'],first)
    csv_roundtrip(out/'email_repeat.csv',['Fajl','Targy','Celmappa'],second)
    assert subject('From: x\r\nSUBJECT: SZÁMLA és rendelés\r\n\r\nbody')=='számla és rendelés'
    assert category('SZÁMLA és rendelés')=='Szamlazas'
    assert category('rendeles panasz')=='Rendelesek'
    assert category(subject('From: x\n\nbody'))=='maradt'
    assert category('Fw: PANASZ')=='Panaszok'
    # Collision source/target survive with distinct bytes.
    src=folder/'collision.eml'; dest=folder/'Szamlazas/collision.eml'
    src.write_bytes(b'new'); dest.write_bytes(b'old')
    try: move_exclusive(src,dest)
    except FileExistsError: pass
    else: raise AssertionError('Email collision overwrote')
    assert src.read_bytes()==b'new' and dest.read_bytes()==b'old'
    return {'first_run':first,'repeat_rows':5,'classification_boundaries':'pass','collision':'no overwrite',
            'body_not_classified':'Intentional source limitation; email07 remains for human review',
            'native_uipath_or_excel_log_run':False}


def local08web(out):
    from lxml import html
    doc=html.parse(str(ROOT/'nap08/nap08/docs/webshop_admin.html'))
    header=[s.strip() for s in doc.xpath('//table//th/text()')]
    records=[[c.text_content().strip() for c in tr.xpath('./td')] for tr in doc.xpath('//table//tr[td]')]
    assert header==['Rendelésszám','Vevő','Város','Összeg','Státusz']
    assert len(records)==25 and all(len(r)==5 for r in records)
    assert records[0][0]=='R-79001' and records[-1][0]=='R-79025'
    new=[r for r in records if r[4]=='új']
    assert [r[0] for r in new]==['R-79003','R-79006','R-79011','R-79015','R-79020','R-79024']
    csv_roundtrip(out/'Rendelesek.csv',header,records)
    csv_roundtrip(out/'Teendok.csv',header,new)
    return {'source_rows':25,'columns':header,'new_ids':[r[0] for r in new],
            'native_browser_table_extraction':False,'csv_export_roundtrip':True}


def local08errors(out):
    from lxml import html
    from pypdf.errors import PdfReadError, PdfStreamError
    from io import BytesIO
    from pypdf import PdfReader
    bad=ROOT/'nap08/nap08/docs/invoice_X_serult.pdf'
    try: invoice_fields(bad)
    except ValueError as exc: assert str(exc)=='business:no_readable_text'
    else: raise AssertionError('Blank PDF silently became success')
    try: PdfReader(BytesIO(b'not a PDF'))
    except (PdfReadError,PdfStreamError): pass
    else: raise AssertionError('Technical corruption undetected')
    invoice_result=json.loads((out/'pdf-results.json').read_text(encoding='utf-8'))
    good=sorted(Path(invoice_result['results']['source_directory']).rglob('*.pdf'))
    assert len(good)==11
    mixed=good[:5]+[bad]+good[5:]
    success=[]; errors=[]
    for path in mixed:
        try: success.append(invoice_fields(path))
        except ValueError as exc: errors.append([path.name,str(exc)])
    assert len(success)==11 and len(errors)==1 and all(all(r) for r in success)
    log=out/'hibak.txt'
    line='Hibás PDF: invoice_X_serult.pdf | A PDF nem tartalmaz kiolvasható szöveget.\n'
    for count in (1,2):
        with log.open('a',encoding='utf-8') as stream: stream.write(line)
        assert len(log.read_text(encoding='utf-8').splitlines())==count
    raw=(ROOT/'nap08/nap08/docs/gyakorlo_karportal.html').read_text(encoding='utf-8')
    assert 'btn-' in raw and ('Math.random' in raw or 'randomUUID' in raw)
    doc=html.fromstring(raw)
    buttons=doc.xpath('//button[normalize-space(.)="Piszkozat mentése"]')
    assert len(buttons)==1
    # Synthetic reload, never execute source JavaScript or operate live portal.
    button=buttons[0]; button.set('id','btn-first')
    assert len(doc.xpath('//button[@id="btn-first"]'))==1
    button.set('id','btn-second')
    assert not doc.xpath('//button[@id="btn-first"]')
    assert len(doc.xpath('//button[normalize-space(.)="Piszkozat mentése"]'))==1
    return {'provided_bad_pdf':'valid PDF with no extractable text, business branch',
            'synthetic_corrupt_pdf':'technical branch','mixed_batch':{'processed':12,'success':11,'errors':1},
            'append_history_lines':2,'stable_selector_offline':'unique across synthetic ID change',
            'live_portal_execution':False,'portal_constraint':'user-started only'}


def verify_excel_robot(out):
    """Verify previously observed UiPath runs; this probe never starts a robot."""
    project = ROOT/'.bpa/projects/ExcelRobot'
    supplied = ROOT/'nap07/nap07/docs/ugyfelek.xlsx'
    assert digest(project/'input/ugyfelek.xlsx') == digest(supplied)
    original = rows(supplied, 'Ugyfelek')
    results = []
    for suffix, threshold, highlighted, filename in (
        ('main', 100000, 6, 'ugyfelek_minositve.xlsx'),
        ('150000', 150000, 3, 'ugyfelek_minositve_v2.xlsx'),
    ):
        transcript = out/('uipath-run-'+suffix+'.json')
        raw = transcript.read_text(encoding='utf-8-sig')
        # The official CLI emits log lines before its JSON envelope.
        envelope = json.loads(raw[raw.index('{'):])
        assert envelope['Result'] == 'Success'
        assert envelope['Data']['hasErrors'] is False
        assert 'rows=12' in raw and 'execution ended in:' in raw
        output = project/'output'/filename
        actual = rows(output, 'Ugyfelek')
        assert len(actual) == 13 and actual[0] == original[0]
        for source_row, actual_row in zip(original[1:], actual[1:]):
            assert actual_row[:3] == source_row[:3]
            assert actual_row[3] == ('kiemelt' if source_row[2] >= threshold else 'normál')
        assert sum(row[3] == 'kiemelt' for row in actual[1:]) == highlighted
        results.append({'threshold':threshold, 'rows':12, 'highlighted':highlighted,
                        'workbook':str(output.relative_to(ROOT)), 'sha256':digest(output),
                        'runtime_transcript':transcript.name})
    assert (out/'excel-native.jpg').is_file() and (out/'excel-native.txt').is_file()
    return {'current_nap07_input_byte_identical':True, 'verified_actual_runs':results,
            'execution_transport':'official direct CLI; MCP build passed separately',
            'native_excel_open_observed':True, 'probe_starts_robot':False}


PHASES={'local03':local03,'live':live,'local07':local07,'pdf':local08pdf,'email':local08email,'web':local08web,'errors':local08errors,'verify-excel-robot':verify_excel_robot}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('phase',choices=PHASES)
    parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args()
    out=args.out.resolve()
    if not out.is_relative_to(ROOT/'.bpa'):
        parser.error('Diagnostic output must stay inside this checkout .bpa directory')
    out.mkdir(parents=True,exist_ok=True)
    record={'phase':args.phase,'checked_utc':datetime.now(timezone.utc).isoformat(),'robot_execution':False}
    try:
        record['results']=PHASES[args.phase](out)
        record['status']='pass'
    except Exception as exc:
        record['status']='failed'
        # Never stringify live exceptions with potentially sensitive server data.
        record['error_type']=type(exc).__name__
        if args.phase!='live': record['detail']=str(exc)
        save(out/(args.phase+'-results.json'),record)
        print(json.dumps(record,ensure_ascii=True))
        return 1
    save(out/(args.phase+'-results.json'),record)
    print(json.dumps({'phase':args.phase,'status':record['status'],'evidence':str(out/(args.phase+'-results.json'))}))
    return 0


if __name__=='__main__':
    sys.exit(main())

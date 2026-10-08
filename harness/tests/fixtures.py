"""Generate minimal synthetic assets; no professor bundles or real course data."""
from pathlib import Path
import zipfile
from openpyxl import Workbook
from pptx import Presentation


def make_lesson_inputs(root):
    docs = Path(root) / 'nap04/nap04/docs'
    docs.mkdir(parents=True, exist_ok=True)
    for filename, sheetname in [('ugyfelek.xlsx', 'Sheet1'), ('Raw_data.xlsx', 'us-500'),
                                ('karbejelentesek_input.xlsx', 'bejelentesek')]:
        book = Workbook()
        sheet = book.active; sheet.title = sheetname
        if filename.startswith('karbejelentesek'):
            sheet.append(['kotvenyszam', 'kartipus', 'karesemeny_datum', 'becsult_osszeg', 'karleiras', 'ugyszam'])
            for n in range(10):
                sheet.append(['TR-402318' if n == 0 else f'TR-{800000+n}', 'vízkár', '2025-11-03', 1000+n, 'Synthetic fixture', None])
        else:
            sheet.append(['Name', 'Amount'])
            for n in range(12):
                sheet.append([f'Synthetic {n}', 1000+n])
        book.save(docs / filename); book.close()
    deck = Presentation()
    for n in range(9):
        deck.slides.add_slide(deck.slide_layouts[1])
    deck.save(docs / 'Dickinson_Sample_Slides.pptx')
    (docs / 'gyakorlo_karportal.html').write_text('<h1>SYNTHETIC PREPARATION FIXTURE; no interactive portal</h1>', encoding='utf-8')
    with zipfile.ZipFile(docs / 'gyakorlo_mappa.zip', 'w') as archive:
        for ext in ('pdf', 'xlsx', 'docx'):
            for n in range(5):
                archive.writestr(f'gyakorlo_mappa/synthetic{n}.{ext}', b'SYNTHETIC PREPARATION FIXTURE')
        for n in range(3):
            archive.writestr(f'gyakorlo_mappa/synthetic{n}.png', b'SYNTHETIC PREPARATION FIXTURE')

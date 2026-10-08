"""Behavioral checks in isolated workspaces; no service credentials or robot run."""
import io
import json
from pathlib import Path
import shutil
import tempfile
import unittest
import zipfile
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation

from harness.cli import ROOT, initialize, transition, add_task, handoff
from harness.projects import AUTHOR_ROOT, prepare_project
from harness.audit_share import inspect_bytes


class PortableTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='BPA different laptop with spaces ')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        for name in ['TASK_PROGRESS.md', '.env.example', 'harness/config.example.json'] + [f'nap0{n}/notes/TASK_PROGRESS.md' for n in range(1, 5)]:
            dst = self.root / name
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(ROOT / name, dst)
        initialize(self.root)

    def data(self):
        return json.loads((self.root / '.bpa/progress.json').read_text(encoding='utf-8'))

    def test_fresh_progress_has_no_historical_success_or_owner(self):
        data = self.data()
        self.assertGreater(len(data['tasks']), 55)
        for task in data['tasks']:
            self.assertFalse(task['owner'])
            self.assertNotIn('Status: `done`', task['requirements'])
            for step in task['steps']:
                self.assertEqual(step['state'], 'not_started')
                self.assertFalse(step['evidence'])

    def test_initialization_preserves_credentials_and_existing_progress(self):
        env = self.root / '.env.local'
        env.write_text('CMC_PRO_API_KEY=synthetic-local-value\n', encoding='utf-8')
        before = (self.root / '.bpa/progress.json').read_bytes()
        initialize(self.root)
        self.assertEqual(before, (self.root / '.bpa/progress.json').read_bytes())
        self.assertIn('synthetic-local-value', env.read_text())

    def test_claim_order_evidence_and_owner_are_enforced(self):
        task = next(t for t in self.data()['tasks'] if len(t['steps']) > 1)['id']
        for number, state, evidence, owner, checked in [
                (2, 'in_progress', 'read', 'chat-a', True),
                (1, 'done', 'result', 'chat-a', True),
                (1, 'in_progress', '', 'chat-a', True),
                (1, 'in_progress', 'read', 'chat-a', False)]:
            with self.assertRaises(ValueError):
                transition(self.root, task, number, state, evidence, owner, checked)
        transition(self.root, task, 1, 'in_progress', 'Read lesson and checked dependencies', 'chat-a', True)
        with self.assertRaises(ValueError):
            transition(self.root, task, 1, 'done', 'Actual check passed', 'chat-b')
        transition(self.root, task, 1, 'blocked', 'Precise missing private sign-in', 'chat-a')
        with self.assertRaises(ValueError):
            transition(self.root, task, 2, 'in_progress', 'read', 'chat-a', True)
        transition(self.root, task, 1, 'in_progress', 'Sign-in provided; resume', 'chat-a', True)
        transition(self.root, task, 1, 'done', 'Actual evidence path', 'chat-a')
        with self.assertRaises(ValueError):
            transition(self.root, task, 1, 'in_progress', 'Overwrite old evidence', 'chat-a', True)
        transition(self.root, task, 2, 'in_progress', 'Dependencies checked', 'chat-a', True)
        self.assertEqual(len((self.root / '.bpa/events.jsonl').read_text().splitlines()), 5)

    def copy_sources(self, name):
        sources = {'Proba': 'nap01/deliverables/Proba',
                   'PortalRobot': 'nap04/deliverables/PortalRobot/setup_2026-10-06/BPA_Setup_Smoke',
                   'CMC_auto_refresh': 'nap03/deliverables/CMC_auto_refresh'}
        source = sources.get(name, 'nap04/deliverables/' + name)
        shutil.copytree(ROOT / source, self.root / source,
                        ignore=shutil.ignore_patterns('.local', '.project', 'node_modules'), dirs_exist_ok=True)
        shutil.copytree(ROOT / 'nap04/nap04/docs', self.root / 'nap04/nap04/docs', dirs_exist_ok=True)
        if name == 'CMC_auto_refresh':
            dest = self.root / 'nap03/deliverables/CMC_arfolyamok_SOL.xlsx'
            shutil.copyfile(ROOT / dest.relative_to(self.root), dest)

    def test_new_task_and_explicit_handoff_preserve_evidence(self):
        add_task(self.root, {'id': 'LOCAL-01', 'title': 'New requirement', 'steps': ['Inspect inputs', 'Verify output']})
        transition(self.root, 'LOCAL-01', 1, 'in_progress', 'Read inputs', 'chat-a', True)
        with self.assertRaises(ValueError):
            handoff(self.root, 'LOCAL-01', 'chat-b', 'chat-c', 'Resume step1')
        handoff(self.root, 'LOCAL-01', 'chat-a', 'chat-b', 'Resume step1; actual check pending')
        transition(self.root, 'LOCAL-01', 1, 'done', 'Actual inspected inputs', 'chat-b')
        with self.assertRaises(ValueError):
            add_task(self.root, {'id': 'LOCAL-01', 'title': 'Duplicate', 'steps': ['Overwrite']})

    def test_simultaneous_update_cannot_overwrite_another_claim(self):
        lock = self.root / '.bpa/progress.lock'
        lock.write_text('another updater', encoding='utf-8')
        before = (self.root / '.bpa/progress.json').read_bytes()
        with self.assertRaises(ValueError):
            add_task(self.root, {'id': 'LOCAL-02', 'title': 'Concurrent', 'steps': ['Inspect']})
        self.assertEqual(before, (self.root / '.bpa/progress.json').read_bytes())
        self.assertTrue(lock.exists())

    def test_all_projects_relocate_and_preserve_reference_bytes(self):
        for name in ['Proba', 'Nap04_gyakorlas', 'ExcelRobot', 'RiportRobot', 'Fajlrendezo', 'PortalRobot', 'CMC_auto_refresh']:
            with self.subTest(name=name):
                self.copy_sources(name)
                target = prepare_project(self.root, name)
                project = json.loads((target / 'project.json').read_text())
                self.assertEqual(project['targetFramework'], 'Windows')
                self.assertEqual(project['expressionLanguage'], 'VisualBasic')
                for path in target.glob('*.xaml'):
                    etree.parse(str(path))
                    self.assertNotIn(AUTHOR_ROOT, path.read_text())
                    self.assertNotIn(AUTHOR_ROOT.replace('/', '\\'), path.read_text())
                self.assertFalse(list(target.rglob('*.lock')))
                self.assertFalse(list(target.rglob('.local')))
                with self.assertRaises(ValueError):
                    prepare_project(self.root, name)

    def test_report_has_fresh_data_and_no_old_chart_or_run_marker(self):
        self.copy_sources('RiportRobot')
        target = prepare_project(self.root, 'RiportRobot')
        for n in (12, 6):
            folder = target / f'reports_{n}'
            self.assertEqual((folder / 'Riport_adatok.xlsx').read_bytes(), (self.root / 'nap04/nap04/docs/Raw_data.xlsx').read_bytes())
            slides = Presentation(folder / 'Riport_eredmeny.pptx').slides
            self.assertGreaterEqual(len(slides), 9)
            self.assertFalse(any(s.has_chart or s.name.startswith('BPA_') for s in slides[8].shapes))
            self.assertTrue(any(s.name == 'Content Placeholder 2' for s in slides[8].shapes))

    def test_portal_stays_closed_and_has_fresh_exact_input(self):
        self.copy_sources('PortalRobot')
        target = prepare_project(self.root, 'PortalRobot')
        self.assertFalse((target / 'Portal_Batch_Resume.xaml').exists())
        batch = etree.parse(str(target / 'REFERENCE_ONLY_Portal_Batch_Resume.xaml'))
        var = batch.xpath('//*[local-name()="Variable"][@Name="batchBaselineReviewed"]')[0]
        self.assertEqual(var.get('Default'), 'False')
        self.assertTrue(etree.parse(str(target / 'Main.xaml')).xpath('//*[local-name()="Throw"]'))
        book = load_workbook(target / 'input/karbejelentesek_input.xlsx', data_only=True)
        rows = list(book['bejelentesek'].values)
        self.assertEqual(len(rows), 11)
        self.assertTrue(all(row[5] is None for row in rows[1:]))
        book.close()

    def test_sorter_starts_unsorted_and_keeps_docs(self):
        self.copy_sources('Fajlrendezo')
        target = prepare_project(self.root, 'Fajlrendezo')
        files = list((target / 'exercise/gyakorlo_mappa').glob('*'))
        self.assertEqual(sum(p.suffix.lower() in ('.pdf', '.xlsx', '.docx') for p in files), 15)
        self.assertEqual(sum(p.suffix.lower() == '.png' for p in files), 3)

    def test_secret_scan_inspects_nested_archives_without_printing_value(self):
        credential = b'synthetic-private-test-secret'
        inner = io.BytesIO()
        with zipfile.ZipFile(inner, 'w') as z:
            z.writestr('sheet.xml', credential)
        outer = io.BytesIO()
        with zipfile.ZipFile(outer, 'w') as z:
            z.writestr('workbook.xlsx', inner.getvalue())
        findings = inspect_bytes('fixture.zip', outer.getvalue(), [credential])
        self.assertTrue(any(f['path'].endswith('fixture.zip::workbook.xlsx::sheet.xml') for f in findings))
        self.assertNotIn(credential.decode(), json.dumps(findings))
        self.assertFalse(inspect_bytes('public.txt', b'public quote', [credential]))


if __name__ == '__main__':
    unittest.main()

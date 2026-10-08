"""Prepare fresh, isolated runnable copies; never execute a robot or a browser."""
from pathlib import Path
import json
import re
import shutil
import sys
import zipfile
from xml.sax.saxutils import escape

AUTHOR_ROOT = 'C:/Users/nagya/dev/school/BPA'


def copy_file(source, target):
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, target)


def rebase(text, root, target):
    # Historical paths become new work locations, never the reference output dirs.
    aliases = {
        AUTHOR_ROOT + '/nap04/working/setup_2026-10-06/BPA_Setup_Smoke': target.as_posix(),
        AUTHOR_ROOT + '/nap04/working/setup_2026-10-06': (target / 'evidence').as_posix(),
        AUTHOR_ROOT + '/nap04/working/readiness_2026-10-06': (target / 'input').as_posix(),
        AUTHOR_ROOT + '/nap03/working/cmc_auto_refresh/CMC_Auto_Refresh': target.as_posix(),
        AUTHOR_ROOT: root.as_posix(),
    }
    # Both XAML literal paths and JSON-escaped backslashes occur in reference files.
    for old, new in aliases.items():
        for a, b in ((old, new), (old.replace('/', '\\'), new.replace('/', '\\')),
                     (old.replace('/', '\\\\'), new.replace('/', '\\\\'))):
            text = text.replace(a, escape(b, {'"': '&quot;'}))
    return text


def prepare_project(root, name):
    root = Path(root).resolve()
    target = root / '.bpa/projects' / name
    if target.exists():
        raise ValueError('Work copy already exists; preserve its evidence. Use a new checkout for a full rerun.')
    sources = {
        'Proba': root / 'nap01/deliverables/Proba',
        'PortalRobot': root / 'nap04/deliverables/PortalRobot/setup_2026-10-06/BPA_Setup_Smoke',
        'CMC_auto_refresh': root / 'nap03/deliverables/CMC_auto_refresh',
    }
    source = sources.get(name, root / 'nap04/deliverables' / name)
    if not (source / 'project.json').is_file():
        source = root / 'harness/project_templates' / name
    if not (source / 'project.json').is_file():
        raise ValueError('Reference project missing.')
    required = {'ExcelRobot': ['nap04/nap04/docs/ugyfelek.xlsx'],
                'RiportRobot': ['nap04/nap04/docs/Raw_data.xlsx', 'nap04/nap04/docs/Dickinson_Sample_Slides.pptx'],
                'Fajlrendezo': ['nap04/nap04/docs/gyakorlo_mappa.zip'],
                'PortalRobot': ['nap04/nap04/docs/karbejelentesek_input.xlsx', 'nap04/nap04/docs/gyakorlo_karportal.html'],
                'CMC_auto_refresh': ['nap03/deliverables/CMC_arfolyamok_SOL.xlsx']}.get(name, [])
    missing = [p for p in required if not (root / p).is_file()]
    if missing:
        raise ValueError('Local lesson inputs/workbook required before preparation: ' + ', '.join(missing) + '. Import your own lesson ZIPs or complete the prerequisite workbook task; no partial project was created.')
    target.mkdir(parents=True)
    for path in source.glob('*'):
        if path.suffix in ('.xaml', '.uiproj') or path.name in ('project.json', 'entry-points.json'):
            raw = path.read_text(encoding='utf-8-sig')
            if path.suffix == '.xaml':
                raw = rebase(raw, root, target)
            (target / path.name).write_text(raw, encoding='utf-8')
    for resource in ('.objects', '.screenshots'):
        if (source / resource).is_dir():
            shutil.copytree(source / resource, target / resource)
    notes = ['Prepared offline only. Restore packages, validate in Studio, then perform actual acceptance checks.',
             'Reference run success is not a successful run on this laptop.']
    if name == 'ExcelRobot':
        copy_file(root / 'nap04/nap04/docs/ugyfelek.xlsx', target / 'input/ugyfelek.xlsx')
        (target / 'output').mkdir()
    elif name == 'RiportRobot':
        from openpyxl import load_workbook
        from pptx import Presentation
        for count in (12, 6):
            folder = target / f'reports_{count}'
            copy_file(root / 'nap04/nap04/docs/Raw_data.xlsx', folder / 'Riport_adatok.xlsx')
            # Course data copied from source, no chart/output or completed-run marker.
            # Use the previously verified slide9 layout/placeholder from reference; keep its other eight source slides.
            reference_deck = source / f'reports_{count}/Riport_eredmeny.pptx'
            template = Presentation(reference_deck if reference_deck.is_file() else root / 'nap04/nap04/docs/Dickinson_Sample_Slides.pptx')
            clean_slide = template.slides[8]
            for shape in list(clean_slide.shapes):
                if shape.has_chart or shape.name.startswith('BPA_'):
                    shape._element.getparent().remove(shape._element)
            # The source slide has no content shape. Restore one from its own layout.
            if not any(s.name == 'Content Placeholder 2' for s in clean_slide.shapes):
                layout_content = next(p for p in clean_slide.slide_layout.placeholders
                                      if p.placeholder_format.type == 7)
                clean_slide.shapes.clone_placeholder(layout_content)
                clean_slide.shapes[-1].name = 'Content Placeholder 2'
            template.save(folder / 'Riport_eredmeny.pptx')
            book = load_workbook(folder / 'Riport_adatok.xlsx')
            if 'us-500' not in book:
                raise ValueError('Source report sheet is missing.')
            book.close()
        notes += ['Fresh source workbooks; copied presentation layout with prior chart removed. Native slide9 placeholder must be checked in Studio.']
    elif name == 'Fajlrendezo':
        with zipfile.ZipFile(root / 'nap04/nap04/docs/gyakorlo_mappa.zip') as archive:
            destination = target / 'exercise'
            for item in archive.infolist():
                out = (destination / item.filename).resolve()
                if not out.is_relative_to(destination.resolve()):
                    raise ValueError('Unsafe path in lesson ZIP.')
            archive.extractall(destination)
        # Optional PNG samples are explicit fixtures; sorting must happen in the robot.
        # The supplied ZIP already contains the three PNG mini-exercise files.
    elif name == 'PortalRobot':
        copy_file(root / 'nap04/nap04/docs/karbejelentesek_input.xlsx', target / 'input/karbejelentesek_input.xlsx')
        copy_file(root / 'nap04/nap04/docs/gyakorlo_karportal.html', target / 'input/gyakorlo_karportal.html')
        for path in target.glob('Portal_*.xaml'):
            raw = path.read_text(encoding='utf-8')
            raw = raw.replace('../../readiness_2026-10-06/karbejelentesek_input.xlsx', 'input/karbejelentesek_input.xlsx')
            raw = re.sub(r'(<Variable\b[^>]*Default=")True("[^>]*Name="batchBaselineReviewed")', r'\g<1>False\2', raw)
            # Keep the baseline guard closed; a fresh batch must not use the author's hard-coded skip/case proof.
            if path.name == 'Portal_Batch_Resume.xaml':
                path = path.with_name('REFERENCE_ONLY_Portal_Batch_Resume.xaml')
                (target / 'Portal_Batch_Resume.xaml').unlink()
            path.write_text(raw, encoding='utf-8')
        # Hard stop before any browser activities until the local user captures/validates targets.
        main = target / 'Main.xaml'
        raw = main.read_text(encoding='utf-8')
        from lxml import etree
        tree = etree.fromstring(raw.encode('utf-8'))
        ns = tree.nsmap[None]
        sequence = tree.find('{' + ns + '}Sequence')
        stop = etree.Element('{' + ns + '}Throw', Exception='[New System.InvalidOperationException("Read LOCAL_RUN.md; configure fresh targets and baseline before user-started execution.")]')
        index = next((i for i, node in enumerate(sequence) if '.' not in etree.QName(node).localname), len(sequence))
        sequence.insert(index, stop)
        main.write_bytes(etree.tostring(tree, encoding='utf-8', pretty_print=True))
        notes += ['USER-STARTED ONLY. Main stops safely. Historical batch is REFERENCE_ONLY and its review guard is false.',
                  'User captures current browser/fields/buttons in Studio and opens the copied input portal privately.',
                  'Agent authors/logs/verifies offline. Read harness/WORKFLOWS.md for baseline/trial/batch reconciliation.',
                  'Create a fresh batch after the verified trial; never claim the author case ID as your own.']
    elif name == 'CMC_auto_refresh':
        copy_file(root / 'nap03/deliverables/CMC_arfolyamok_SOL.xlsx', target / 'CMC_Auto_Refresh.xlsx')
        path = target / 'Main.xaml'
        raw = path.read_text(encoding='utf-8')
        raw = raw.replace('nap03/working/fetch_cmc_power_query.py', 'harness/course_tools/day03/fetch_cmc_power_query.py')
        python_exe = root / '.venv/Scripts/python.exe'
        raw = re.sub(r'System.Diagnostics.ProcessStartInfo\("[^"]*python.exe"\)',
                     lambda _: 'System.Diagnostics.ProcessStartInfo("' + escape(str(python_exe)) + '")', raw)
        marker = 'info.ArgumentList.Add("--solana")'
        raw = raw.replace(marker, marker + '&#10;info.ArgumentList.Add("--output-dir")&#10;info.ArgumentList.Add("' + escape(target.as_posix() + '/quotes') + '")')
        path.write_text(raw, encoding='utf-8')
        (target / 'quotes').mkdir()
        notes += ['Run harness/configure-cmc.ps1 on this fresh copy before Studio. It rebases native workbook queries without reading credentials.',
                  'Actual robot fetches using the root .env.local, then native Excel Power Query refreshes public JSON.']
    (target / 'LOCAL_RUN.md').write_text('\n\n'.join(notes) + '\n', encoding='utf-8')
    (target / 'preparation.json').write_text(json.dumps({'project': name, 'scope': 'offline preparation only',
        'source': source.relative_to(root).as_posix(), 'actual_run': False}, indent=2), encoding='utf-8')
    return target

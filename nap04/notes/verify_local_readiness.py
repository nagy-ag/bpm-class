"""Read-only source inspection and isolated file/copy/archive checks, not robot runs."""
import csv
import hashlib
import json
import shutil
import zipfile
from collections import Counter
from pathlib import Path
import openpyxl
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'nap04/nap04/docs'
WORK = ROOT / 'nap04/working/readiness_2026-10-06'
WORK.mkdir(parents=True, exist_ok=True)
report = {'source_files_unchanged': True, 'robot_execution': False}
names = ['ugyfelek.xlsx', 'Raw_data.xlsx', 'karbejelentesek_input.xlsx',
         'Dickinson_Sample_Slides.pptx', 'gyakorlo_karportal.html']
hashes = {name: hashlib.sha256((SOURCE/name).read_bytes()).hexdigest() for name in names}
for name in names:
    target = WORK/name
    if not target.exists():
        shutil.copy2(SOURCE/name, target)
    assert hashlib.sha256(target.read_bytes()).hexdigest() == hashes[name]

for name, sheet, headers, rows in [
    ('ugyfelek.xlsx', 'Ugyfelek', ['Nev','Varos','EvesDij','Minosites'], 12),
    ('Raw_data.xlsx', 'us-500', ['State','Quantity','UnitPrice','Total'], 243),
    ('karbejelentesek_input.xlsx', 'bejelentesek',
     ['kotvenyszam','kartipus','karesemeny_datum','becsult_osszeg','karleiras','ugyszam'], 10)]:
    wb = openpyxl.load_workbook(WORK/name, read_only=True, data_only=True)
    table = list(wb[sheet].values)
    assert list(table[0]) == headers and len(table)-1 == rows
    report[name] = {'headers_verified': True, 'data_rows': rows}
    if name == 'ugyfelek.xlsx':
        amounts = [r[2] for r in table[1:]]
        assert sum(x >= 100000 for x in amounts) == 6
        assert sum(x >= 150000 for x in amounts) == 3
        report['classification_fixtures'] = {'threshold100000': 6, 'threshold150000': 3}
    wb.close()
deck = Presentation(WORK/'Dickinson_Sample_Slides.pptx')
assert len(deck.slides) == 9
report['presentation'] = {'slides': 9, 'readable_editable_package': True}

quote_path = ROOT/'nap03/working/readiness_2026-10-06/quotes.json'
quotes = json.loads(quote_path.read_text(encoding='utf-8'))['coins']
with (WORK/'quotes_roundtrip.csv').open('w', encoding='utf-8', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(quotes[0]))
    writer.writeheader()
    writer.writerows(quotes)
with (WORK/'quotes_roundtrip.csv').open(encoding='utf-8', newline='') as f:
    saved = list(csv.DictReader(f))
assert [r['symbol'] for r in saved] == ['BTC','ETH','SOL']
assert all(float(row['price_eur']) == quotes[i]['price_eur'] for i,row in enumerate(saved))
report['json_csv_roundtrip'] = True

sort_root = WORK/'isolated_sort_test'
first_extraction = not sort_root.exists()
sort_root.mkdir(exist_ok=True)
with zipfile.ZipFile(SOURCE/'gyakorlo_mappa.zip') as archive:
    archive_hashes = {Path(info.filename).name: hashlib.sha256(archive.read(info)).hexdigest()
                      for info in archive.infolist() if not info.is_dir()}
    for info in archive.infolist():
        destination = (sort_root/info.filename).resolve()
        assert destination.is_relative_to(sort_root.resolve())
        if first_extraction:
            archive.extract(info, sort_root)
files = [p for p in sort_root.rglob('*') if p.is_file()]
assert len(files) == 20
before = {p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
assert before == archive_hashes
categories = {'.pdf':'PDF','.xlsx':'Excel','.docx':'Word','.png':'Kepek'}
for p in files:
    category = categories.get(p.suffix.lower())
    if category:
        dest_dir = sort_root/category
        dest_dir.mkdir(exist_ok=True)
        destination = (dest_dir/p.name).resolve()
        assert p.resolve().is_relative_to(sort_root.resolve()) and destination.is_relative_to(sort_root.resolve())
        if p.resolve() == destination:
            continue
        assert not destination.exists()
        p.rename(destination)
after = {p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sort_root.rglob('*') if p.is_file()}
assert before == after
counts = Counter(p.parent.name if p.parent.name in categories.values() else 'root'
                 for p in sort_root.rglob('*') if p.is_file())
assert counts == {'PDF':6,'Excel':4,'Word':5,'Kepek':3,'root':2}
report['isolated_file_moves'] = {'counts':dict(counts), 'all20_hashes_preserved':True}
assert all(hashlib.sha256((SOURCE/name).read_bytes()).hexdigest() == hashes[name] for name in names)
(WORK/'local_checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
